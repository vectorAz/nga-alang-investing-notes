#!/usr/bin/env python3
"""Export per-author reply archives to a multi-sheet Excel workbook.

Each author gets one sheet, sorted by post time (ascending) then PID, so the
sheet order can be reconciled line by line against references/originals/.
The originals backup (references/originals/, driven by manifest.json) is the
preferred source; references/sources/*.jsonl is only a fallback for older
checkouts that have no originals backup yet.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
ORIGINALS = ROOT / "references" / "originals"
SOURCES = ROOT / "references" / "sources"
PRIMARY_TID = 47288722

DEFAULT_USERS = [
    (150058, "阿狼", "楼主"),
    (67142257, "兰斯洛t", "重点用户"),
    (60916468, "灰兔尾", "重点用户"),
    (67145714, "Plezl", "重点用户"),
    (39700287, "PoW", "重点用户"),
    (65329649, "65329649", "重点用户"),
]

COLUMNS = [
    ("序号", 6), ("发送时间", 17), ("用户", 11), ("UID", 11), ("PID", 12),
    ("作者分页位", 11), ("字数", 7), ("含引用", 8), ("正文", 92), ("TID", 11),
    ("主题", 36), ("来源", 10), ("原帖链接", 50),
]

SOURCE_LABELS = {"nga": "NGA抓取", "excel-archive": "Excel归档"}


def record_source(record: dict) -> str:
    return SOURCE_LABELS.get(record.get("source") or "", record.get("source") or "未知")


def load_rows() -> tuple[list[dict], str]:
    """Return (records, source label). Prefers the originals backup over sources."""
    manifest_path = ORIGINALS / "manifest.json"
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        rows = []
        for item in manifest["files"]:
            path = ORIGINALS / item["path"]
            with path.open(encoding="utf-8") as handle:
                for line in handle:
                    line = line.strip()
                    if line:
                        rows.append(json.loads(line))
        if rows:
            return rows, f"references/originals（{len(manifest['files'])} 个分片）"

    rows = []
    for path in sorted(SOURCES.glob("*.jsonl")):
        with path.open(encoding="utf-8") as handle:
            for line in handle:
                line = line.strip()
                if line:
                    rows.append(json.loads(line))
    return rows, "references/sources"


def parse_time(value: str) -> datetime:
    try:
        return datetime.strptime(value, "%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return datetime(1900, 1, 1)


def body_text(record: dict) -> str:
    text = (record.get("plain_text") or "").strip()
    if text:
        return text
    if record.get("images") or record.get("attachments"):
        return "[仅图片/附件，无文字]"
    return "[无文字内容]"


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description="按用户分 sheet 导出完整回复")
    parser.add_argument("--out", required=True, help="输出 xlsx 路径")
    parser.add_argument("--uid", type=int, action="append", help="只导出指定 UID，可重复")
    parser.add_argument("--main-only", action="store_true", help="只导出主帖 tid=47288722 的回复")
    args = parser.parse_args()

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("缺少 openpyxl，请先执行: python -m pip install openpyxl", file=sys.stderr)
        return 2

    rows, source_label = load_rows()
    if not rows:
        print(f"未找到数据: {ORIGINALS} 或 {SOURCES}", file=sys.stderr)
        return 1

    users = [u for u in DEFAULT_USERS if not args.uid or u[0] in args.uid]
    if args.uid:
        known = {u[0] for u in DEFAULT_USERS}
        for uid in args.uid:
            if uid not in known:
                users.append((uid, str(uid), "自定义"))

    grouped: dict[int, list[dict]] = {}
    for record in rows:
        if args.main_only and record.get("tid") != PRIMARY_TID:
            continue
        grouped.setdefault(record["authorid"], []).append(record)

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True, size=10, name="微软雅黑")
    body_font = Font(size=10, name="微软雅黑")
    mono_font = Font(size=10, name="Consolas")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    workbook = Workbook()
    overview = workbook.active
    overview.title = "总览"
    overview["A1"] = "NGA 狼楼用户完整回复归档"
    overview["A1"].font = Font(size=14, bold=True, name="微软雅黑", color="1F3864")
    overview["A2"] = f"生成时间：{datetime.now():%Y-%m-%d %H:%M}　　数据源：{source_label}"
    overview["A2"].font = Font(size=9, name="微软雅黑", color="595959")
    overview["A3"] = (
        f"范围：{'仅主帖 tid=' + str(PRIMARY_TID) if args.main_only else '全部已登记主题（含主帖）'}"
    )
    overview["A3"].font = Font(size=9, name="微软雅黑", color="595959")

    headings = ["用户", "UID", "身份", "主帖回复", "其他主题", "合计",
                "NGA抓取", "Excel归档", "起始时间", "结束时间", "平均字数"]
    for index, heading in enumerate(headings, 1):
        cell = overview.cell(row=5, column=index, value=heading)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_cursor = 6
    total = 0
    total_nga = 0
    total_excel = 0
    for uid, name, role in users:
        records = grouped.get(uid, [])
        if not records:
            continue
        times = sorted(parse_time(r["postdate"]) for r in records)
        lengths = [len(body_text(r)) for r in records]
        values = [
            name, str(uid), role,
            sum(1 for r in records if r.get("tid") == PRIMARY_TID),
            sum(1 for r in records if r.get("tid") != PRIMARY_TID),
            len(records),
            sum(1 for r in records if (r.get("source") or "nga") == "nga"),
            sum(1 for r in records if r.get("source") == "excel-archive"),
            times[0].strftime("%Y-%m-%d %H:%M"),
            times[-1].strftime("%Y-%m-%d %H:%M"),
            round(sum(lengths) / len(lengths)),
        ]
        for index, value in enumerate(values, 1):
            cell = overview.cell(row=row_cursor, column=index, value=value)
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if index != 1 else "left", vertical="center")
        total += len(records)
        total_nga += sum(1 for r in records if (r.get("source") or "nga") == "nga")
        total_excel += sum(1 for r in records if r.get("source") == "excel-archive")
        row_cursor += 1

    overview.cell(row=row_cursor, column=1, value="合计").font = Font(size=10, bold=True, name="微软雅黑")
    for index in range(1, 12):
        overview.cell(row=row_cursor, column=index).border = border
    overview.cell(row=row_cursor, column=6, value=total).font = Font(size=10, bold=True, name="微软雅黑")
    overview.cell(row=row_cursor, column=7, value=total_nga).font = Font(size=10, bold=True, name="微软雅黑")
    overview.cell(row=row_cursor, column=8, value=total_excel).font = Font(size=10, bold=True, name="微软雅黑")
    for index in (6, 7, 8):
        overview.cell(row=row_cursor, column=index).alignment = Alignment(horizontal="center")

    notes = [
        "排序规则：Sheet 内按发帖时间升序，同一时间按 PID 升序；序号为该用户的时间顺序位。",
        "「来源」区分记录获取方式：NGA抓取（有 PID，可定位楼层）；Excel归档（历史汇总表导入，无 PID，链接指向主题）。",
        "Excel归档记录的原文出自历史汇总表，单元格内的换行已折叠为空格；「TID」按该表 sheet 内引用块推断。",
        "「作者分页位」是 NGA“只看作者”分页中的位置，不是全帖楼层号；定位请用 PID。",
        "「含引用」为“是”表示回复开头带 [quote]；正文已去除 BBCode。",
        "PID=0 表示主楼本身，PID 为空表示该条来自 Excel 归档、本就没有 PID。",
    ]
    for index, note in enumerate(notes):
        cell = overview.cell(row=row_cursor + 2 + index, column=1, value=note)
        cell.font = Font(size=9, name="微软雅黑", color="404040")

    overview.column_dimensions["A"].width = 26
    for column, width in [("B", 12), ("C", 11), ("D", 11), ("E", 11), ("F", 10),
                          ("G", 10), ("H", 11), ("I", 19), ("J", 19), ("K", 10)]:
        overview.column_dimensions[column].width = width
    overview.freeze_panes = "A6"

    for uid, name, role in users:
        records = sorted(grouped.get(uid, []), key=lambda r: (parse_time(r["postdate"]), r.get("pid") or 0))
        if not records:
            continue
        sheet = workbook.create_sheet(title=name[:31])
        for index, (heading, width) in enumerate(COLUMNS, 1):
            cell = sheet.cell(row=1, column=index, value=heading)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[1].height = 22

        for order, record in enumerate(records, 1):
            row = order + 1
            text = body_text(record)
            quoted = "是" if (record.get("raw_bbcode") or "").lstrip().startswith("[quote]") else ""
            sheet.cell(row=row, column=1, value=order).font = body_font
            time_cell = sheet.cell(row=row, column=2, value=parse_time(record["postdate"]))
            time_cell.number_format = "yyyy-mm-dd hh:mm"
            time_cell.font = mono_font
            sheet.cell(row=row, column=3, value=record.get("author")).font = body_font
            sheet.cell(row=row, column=4, value=record.get("authorid")).font = mono_font
            sheet.cell(row=row, column=5, value=record.get("pid")).font = mono_font
            sheet.cell(row=row, column=6, value=record.get("source_lou")).font = mono_font
            sheet.cell(row=row, column=7, value=len(text)).font = mono_font
            sheet.cell(row=row, column=8, value=quoted).font = body_font
            body_cell = sheet.cell(row=row, column=9, value=text)
            body_cell.font = body_font
            body_cell.alignment = Alignment(vertical="top", wrap_text=False)
            sheet.cell(row=row, column=10, value=record.get("tid")).font = mono_font
            sheet.cell(row=row, column=11, value=record.get("thread_subject")).font = body_font
            sheet.cell(row=row, column=12, value=record_source(record)).font = body_font
            url = record.get("source_url") or (
                f"https://bbs.nga.cn/read.php?pid={record.get('pid')}&opt=128"
                if record.get("pid") is not None else ""
            )
            link_cell = sheet.cell(row=row, column=13, value=url)
            link_cell.font = Font(size=9, name="Consolas", color="0563C1", underline="single")
            if url:
                link_cell.hyperlink = url
            for index in range(1, 14):
                sheet.cell(row=row, column=index).border = border
                if index != 9:
                    sheet.cell(row=row, column=index).alignment = Alignment(
                        horizontal="center" if index in (1, 4, 5, 6, 7, 8, 10, 12) else "left",
                        vertical="top",
                    )
            sheet.row_dimensions[row].height = 15

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:M{len(records) + 1}"
        print(f"  {name}: {len(records)} 行")

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    print(f"已生成 {out_path}（{total} 条记录，{out_path.stat().st_size / 1024 / 1024:.2f} MB）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
