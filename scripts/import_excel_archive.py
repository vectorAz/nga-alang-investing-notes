#!/usr/bin/env python3
"""Import an Excel reply archive into references/originals/.

The historical Excel workbooks only carry two columns (post time, reply body),
so they have neither PID nor thread id. This script therefore:

* infers the thread id per sheet from the tid/pid markers that appear inside the
  quoted blocks (the dominant value per sheet), and
* keys such records as ``excel:<uid>:<yyyymmddHHMM>:<hash>`` so they can never
  collide with ``post:<tid>:<pid>:<uid>`` records scraped from the site.

Rows that already exist in the archive are skipped by matching
(uid, post time) plus normalized body containment, because the older scrape
sometimes stored a truncated quote header for the very same post.

Usage:
    python scripts/import_excel_archive.py --workbook "path/to/replies.xlsx" --dry-run
    python scripts/import_excel_archive.py --workbook "path/to/replies.xlsx" --batch 20260926
"""
from __future__ import annotations

import argparse
import difflib
import hashlib
import html
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from archive_store import load_archive, verify_archive, write_archive  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
ORIGINALS = ROOT / "references" / "originals"
SHANGHAI = timezone(timedelta(hours=8))

# author id -> display name, as stored in the archive
AUTHOR_NAMES = {
    150058: "-阿狼-",
    67142257: "兰斯洛t",
}

# Sheet names that belong to a tracked user other than the default author.
SHEET_AUTHOR_HINTS = [
    ("兰斯洛", 67142257),
]

# Thread titles for the long-running series; unknown ids fall back to a label.
THREAD_TITLES = {
    8869728: "2016-2017 历史主题",
    24900465: "2020-2021 历史主题",
    30081858: "2022 历史主题",
    43098323: "我必不是蛇年红包（2025）",
    45974302: "自立自强，科学技术打头阵",
    47207407: "自立自强，科学技术打头阵(结束使命，封存)",
    47288722: "自立自强，科学技术打头阵----救赎版",
}

PRIMARY_TID = 47288722
LEGACY_OP_TIDS = {8869728, 24900465, 30081858, 43098323, 45974302}

TAG_RE = re.compile(r"\[/?[a-zA-Z][^\[\]]{0,60}\]")


def norm(text: str) -> str:
    """Normalized comparison key: tags, entities and all whitespace removed."""
    value = html.unescape(text or "")
    value = re.sub(r"<[^>]+>", "", value)
    value = TAG_RE.sub("", value)
    value = unicodedata.normalize("NFKC", value)
    return re.sub(r"\s+", "", value)


def to_plain(text: str) -> str:
    """BBCode -> readable plain text (quoted content kept, images annotated)."""
    value = text or ""
    value = re.sub(r"<br\s*/?>", "\n", value, flags=re.I)
    value = re.sub(r"\[img\](.*?)\[/img\]", r"[图片: \1]", value, flags=re.S | re.I)
    value = re.sub(r"\[url=[^\]]*\](.*?)\[/url\]", r"\1", value, flags=re.S | re.I)
    value = re.sub(r"\[url\](.*?)\[/url\]", r"\1", value, flags=re.S | re.I)
    value = re.sub(r"\[s:[^\]]+\]", "", value)
    value = re.sub(r"\[\*\]", "· ", value)
    value = TAG_RE.sub("", value)
    value = re.sub(r"<[^>]+>", "", value)
    value = html.unescape(value)
    value = re.sub(r"[ \t\u3000]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def sheet_author(sheet: str, default_uid: int) -> int:
    for hint, uid in SHEET_AUTHOR_HINTS:
        if hint in sheet:
            return uid
    return default_uid


def infer_tid(texts: list[str]) -> tuple[int, dict]:
    """Pick the thread id that the sheet's quoted blocks mostly reference."""
    counter: Counter[str] = Counter()
    for text in texts:
        if not text:
            continue
        counter.update(re.findall(r"tid=(\d+)", text))
        counter.update(re.findall(r"pid=\d+,(\d+),\d+", text))
    if not counter:
        return 0, {}
    tid, _ = counter.most_common(1)[0]
    return int(tid), dict(counter.most_common(5))


def tracking_role(uid: int, tid: int) -> str:
    if uid == 150058:
        return "op" if tid == PRIMARY_TID else "legacy_op"
    return "priority_user"


def build_record(row: dict, batch: str, workbook_name: str) -> dict:
    uid: int = row["uid"]
    tid: int = row["tid"]
    body: str = row["text"]
    stamp = row["time"]
    moment = datetime.strptime(stamp, "%Y-%m-%d %H:%M").replace(tzinfo=SHANGHAI)
    compact = stamp.replace("-", "").replace(" ", "").replace(":", "")
    digest = hashlib.sha1(norm(body).encode("utf-8")).hexdigest()[:12]
    return {
        "record_key": f"excel:{uid}:{compact}:{digest}",
        "record_kind": "post",
        "author": AUTHOR_NAMES.get(uid, str(uid)),
        "authorid": uid,
        "pid": None,
        "tid": tid,
        "postdate": stamp,
        "postdatetimestamp": int(moment.timestamp()),
        "subject": "",
        "thread_subject": THREAD_TITLES.get(tid, f"历史主题 tid={tid}" if tid else "历史主题（未识别）"),
        "raw_bbcode": body,
        "raw_format": "bbcode",
        "plain_text": to_plain(body),
        "content_length": len(body),
        "source": "excel-archive",
        "source_url": f"https://bbs.nga.cn/read.php?tid={tid}" if tid else "",
        "source_lou": None,
        "tracking_role": tracking_role(uid, tid),
        "score": 0,
        "alterinfo": "",
        "comments": [],
        "hotreply": [],
        "images": [],
        "attachments": [],
        "links": [],
        "import_batch": batch,
        "import_sheet": row["sheet"],
        "import_workbook": workbook_name,
        "text_note": "Excel 导出版本，原文换行已折叠为空格",
    }


def find_duplicate(record: dict, existing_at_time: list[dict]) -> dict | None:
    body = norm(record["raw_bbcode"])
    if not body:
        return None
    for candidate in existing_at_time:
        for field in ("plain_text", "raw_bbcode"):
            other = norm(candidate.get(field) or "")
            if not other:
                continue
            if other == body:
                return candidate
            short, long = sorted((other, body), key=len)
            if len(short) >= 0.75 * len(long) and short in long:
                return candidate
    return None


def find_op_duplicate(record: dict, op_index: dict[int, list[dict]]) -> dict | None:
    """The Excel sheet repeats the thread opener with an older wording and date."""
    body = norm(record["raw_bbcode"])
    if not body:
        return None
    for candidate in op_index.get(record["tid"], []):
        other = norm(candidate.get("plain_text"))
        if not other:
            continue
        if difflib.SequenceMatcher(None, body, other).ratio() >= 0.9:
            return candidate
    return None


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description="把历史 Excel 回复归档并入 originals")
    parser.add_argument("--workbook", required=True, help="xlsx 路径")
    parser.add_argument("--default-uid", type=int, default=150058, help="未命中作者提示时使用的 UID")
    parser.add_argument("--batch", default="", help="批次标签，默认取工作簿名里的日期")
    parser.add_argument("--out", type=Path, help="把待导入记录另存为 JSONL（默认写入 references/_excel_import_<batch>.jsonl）")
    parser.add_argument("--dry-run", action="store_true", help="只统计，不写入归档")
    args = parser.parse_args()

    try:
        from openpyxl import load_workbook
    except ImportError:
        print("缺少 openpyxl，请先执行: python -m pip install openpyxl", file=sys.stderr)
        return 2

    workbook_path = Path(args.workbook).expanduser()
    if not workbook_path.exists():
        print(f"找不到工作簿: {workbook_path}", file=sys.stderr)
        return 1

    batch = args.batch or (re.search(r"(20\d{6})", workbook_path.name).group(1)
                           if re.search(r"(20\d{6})", workbook_path.name) else "excel")

    book = load_workbook(workbook_path, read_only=True, data_only=True)
    raw_rows: list[dict] = []
    per_sheet: dict[str, dict] = {}
    for sheet in book.sheetnames:
        uid = sheet_author(sheet, args.default_uid)
        texts: list[str] = []
        rows: list[dict] = []
        for time_cell, body_cell in book[sheet].iter_rows(min_row=2, values_only=True):
            if time_cell is None and body_cell in (None, ""):
                continue
            stamp = str(time_cell).strip()
            if not re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$", stamp):
                continue
            rows.append({"sheet": sheet, "uid": uid, "time": stamp, "text": str(body_cell) if body_cell else ""})
            texts.append(str(body_cell) if body_cell else "")
        tid, markers = infer_tid(texts)
        for row in rows:
            row["tid"] = tid
        raw_rows.extend(rows)
        per_sheet[sheet] = {"uid": uid, "rows": len(rows), "tid": tid, "markers": markers}
        print(f"  {sheet:26s} uid={uid} {len(rows):6d} 行 → tid={tid} {THREAD_TITLES.get(tid, '')}", file=sys.stderr)

    # 1) 去掉工作簿内部完全重复的行
    seen: set[tuple] = set()
    unique_rows: list[dict] = []
    internal_dupes = 0
    for row in raw_rows:
        key = (row["uid"], row["time"], norm(row["text"]))
        if key in seen:
            internal_dupes += 1
            continue
        seen.add(key)
        unique_rows.append(row)

    # 2) 与现有归档比对
    archive = load_archive(ORIGINALS) if (ORIGINALS / "manifest.json").exists() else []
    by_time: dict[tuple, list[dict]] = defaultdict(list)
    op_index: dict[int, list[dict]] = defaultdict(list)
    for record in archive:
        by_time[(record["authorid"], record.get("postdate"))].append(record)
        if record.get("pid") == 0:
            op_index[record.get("tid")].append(record)

    matched = 0
    matched_op = 0
    matched_by_sheet: Counter = Counter()
    new_records: list[dict] = []
    new_by_sheet: Counter = Counter()
    for row in unique_rows:
        duplicate = find_duplicate({"raw_bbcode": row["text"]}, by_time.get((row["uid"], row["time"]), []))
        if duplicate is not None:
            matched += 1
            matched_by_sheet[row["sheet"]] += 1
            continue
        new_records.append(build_record(row, batch, workbook_path.name))
        new_by_sheet[row["sheet"]] += 1

    # 3) 去掉与现有主楼重复的抓取（同一主题的开帖被 Excel 以旧日期重复收录）
    kept: list[dict] = []
    for record in new_records:
        if find_op_duplicate(record, op_index) is not None:
            matched_op += 1
            new_by_sheet[record["import_sheet"]] -= 1
            continue
        kept.append(record)
    new_records = kept

    # 主键兜底去重（同分钟同内容理论上已被上一步拦截）
    keys = set()
    for record in new_records:
        base = record["record_key"]
        suffix = 1
        while record["record_key"] in keys:
            suffix += 1
            record["record_key"] = f"{base}-{suffix}"
        keys.add(record["record_key"])

    new_records.sort(key=lambda r: (r["postdate"], r["record_key"]))
    out_path = args.out or (ROOT / "references" / f"_excel_import_{batch}.jsonl")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as handle:
        for record in new_records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n")

    by_uid = Counter(r["authorid"] for r in new_records)
    by_year = Counter(r["postdate"][:4] for r in new_records)

    report = {
        "workbook": workbook_path.name,
        "batch": batch,
        "sheets": {name: {"rows": info["rows"], "tid": info["tid"], "uid": info["uid"]} for name, info in per_sheet.items()},
        "rows_read": len(raw_rows),
        "internal_duplicates": internal_dupes,
        "already_in_archive": matched,
        "matched_by_sheet": dict(matched_by_sheet),
        "matched_thread_opener": matched_op,
        "new_records": len(new_records),
        "new_by_sheet": dict(new_by_sheet),
        "new_by_uid": {str(k): v for k, v in by_uid.items()},
        "new_by_year": dict(sorted(by_year.items())),
        "import_jsonl": str(out_path),
        "archive_before": len(archive),
        "archive_after": len(archive) + len(new_records),
    }

    if not args.dry_run:
        write_archive(new_records, ORIGINALS)
        report["verify"] = verify_archive(ORIGINALS)

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
